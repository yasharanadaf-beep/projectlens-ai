from pypdf import PdfReader
from pptx import Presentation


def extract_pdf(file):

    text = ""

    reader = PdfReader(file)

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


def extract_ppt(file):

    text = ""

    presentation = Presentation(file)

    for slide_number, slide in enumerate(
        presentation.slides, start=1
    ):

        text += f"\n--- Slide {slide_number} ---\n"

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                if shape.text.strip():

                    text += shape.text + "\n"

    return text


def extract_text(file):

    filename = file.name.lower()

    if filename.endswith(".pdf"):

        return extract_pdf(file)

    elif filename.endswith(".pptx"):

        return extract_ppt(file)

    else:

        return "Unsupported file type."